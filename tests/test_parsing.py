from datetime import date

from resume_ai.parsing import mask_pii, profile, split_sections, years_from_ranges


def test_sections_are_detected_from_headings():
    s = split_sections("Jane\nSUMMARY\nAnalyst\nWork Experience\n- Built X\nSkills:\nPython\nEducation\nB.Tech")
    assert {"summary", "experience", "skills", "education"} <= set(s)


def test_overlapping_ranges_are_merged():
    text = "Analyst Jan 2020 - Dec 2021\nFreelance Jun 2021 - Mar 2022"
    assert years_from_ranges(text, today=date(2026, 1, 1)) == 2.2


def test_present_ranges_run_to_today():
    assert years_from_ranges("Engineer | Jun 2024 – Present", today=date(2026, 6, 1)) == 2.1


def test_stated_years_used_when_no_dates():
    p = profile("Summary\nData analyst with 4+ years of experience in reporting.\nSkills\nSQL")
    assert p.years_experience == 4.0


def test_education_level_and_contact():
    p = profile("Asha Rao\nasha@example.com | +91 98200 11223\nEducation\nM.Sc. Statistics")
    assert p.education_label == "Master's" and p.has_email and p.has_phone
    assert p.name_guess == "Asha Rao"


def test_pii_is_masked():
    out = mask_pii("Call +91 98200 11223 or mail asha@example.com")
    assert "[phone]" in out and "[email]" in out and "98200" not in out


def test_year_ranges_are_not_phone_numbers():
    assert mask_pii("Analyst 2020-2022") == "Analyst 2020-2022"
    assert not profile("Jane Doe\nExperience\nAnalyst 2019-2023").has_phone


def test_common_date_formats_are_read():
    from datetime import date
    from resume_ai.parsing import years_from_ranges
    today = date(2026, 10, 12)
    cases = {"Jan 2020 – Present": 6.8, "06/2021 - 05/2023": 2.0, "6/2021 to 5/2023": 2.0, "2021-06 to 2023-05": 2.0,
             "Jun'21 – Present": 5.4, "Jun ’21 - May ’23": 2.0, "Jan-2020 to Dec-2022": 3.0, "March 2020 to Till Date": 6.6,
             "Since 2020": 6.8, "Jan 2020 to till now": 6.8, "01.2020 – 12.2022": 3.0, "2019 - 2021": 3.0, "2017 to date": 9.8}
    for text, want in cases.items():
        got = years_from_ranges("Analyst — Acme | " + text, today)
        assert got is not None and abs(got - want) <= 0.25, (text, got)
    assert years_from_ranges("Call +91 98765 43210", today) is None


def test_unusual_headings_and_bullet_symbols():
    from resume_ai.parsing import _heading_key, profile
    assert _heading_key("Professional Experience & Projects") == "experience"
    assert _heading_key("Areas of Expertise") == "skills"
    assert _heading_key("- Data quality training") is None
    assert _heading_key("Python, SQL, Airflow") is None
    p = profile("Work History\n➢ Built Airflow pipelines for 40 teams\n✔ Reduced cost by 20%\nSkill Set\nPython, SQL")
    assert {"experience", "skills"} <= set(p.sections)
    assert p.bullets[0].startswith("Built")


def test_docx_tables_stay_under_their_heading():
    import io
    import docx
    from resume_ai.parsing import extract_text, profile
    d = docx.Document()
    d.add_paragraph("Skills")
    t = d.add_table(rows=1, cols=3)
    for i, s in enumerate(["Python", "SQL", "Airflow"]):
        t.cell(0, i).text = s
    d.add_paragraph("Education")
    d.add_paragraph("B.Tech, 2021")
    buf = io.BytesIO()
    d.save(buf)
    p = profile(extract_text("r.docx", buf.getvalue()))
    assert "Airflow" in p.sections["skills"] and "Airflow" not in p.sections.get("education", "")


def test_non_english_resume_is_flagged():
    from resume_ai.parsing import profile
    p = profile("अर्जुन नायर\nअनुभव\nडेटा इंजीनियर — वुडग्रोव बैंक\nकौशल\nPython")
    assert any("not in English" in w for w in p.warnings)
